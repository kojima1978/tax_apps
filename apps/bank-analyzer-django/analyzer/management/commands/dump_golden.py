"""React 移行の「正解」を記録する（REACT_MIGRATION_PLAN.md 段階1）。

いまの Django の挙動をそのまま JSON に書き出し、新アプリ（apps/bank-analyzer）が
同じ入力で同じ結果を返すかを突き合わせるための材料にする。**挙動を変えるためのものではない**。

    # 架空の入力（リポジトリに入れてよい）
    python manage.py dump_golden synthetic --inputs <dir> --out <dir>
    # 稼働中の案件（公開リポジトリには絶対に入れない。DB の複製に対して走らせる）
    python manage.py dump_golden real --out <dir>

どちらも**使い捨ての DB に対して走らせる前提**。synthetic は案件を作り、real も分類などの
書き込みを試す（書き込みは atomic の中で巻き戻すが、本番 DB には向けないこと）。

再現性のために:
- 取引の id は案件内の id 順の通し番号（1始まり）に置き換える
- それ以外の *_id、作成日時・出力日時は落とす
- ファイル名の日付（R081005）は <DATE> に置き換える
"""
import csv
import io
import json
import math
import re
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction as db_transaction
from django.test import Client
from django.urls import reverse
from openpyxl import load_workbook
from rapidfuzz import fuzz

from analyzer.lib import config, llm_classifier
from analyzer.lib.constants import UNCATEGORIZED
from analyzer.lib.exceptions import CsvImportError
from analyzer.lib.importer import load_csv, validate_balance
from analyzer.lib.text_utils import normalize_text
from analyzer.models import Case, ClassificationChange
from analyzer.services.analysis import AnalysisService
from analyzer.services.transaction import TransactionService

# 取引 id を通し番号に置き換えるキー
TX_ID_KEYS = {"id", "tx_id", "first_tx_id", "transaction_identifier"}
TX_ID_LIST_KEYS = {"tx_ids"}
# 実行のたびに変わるもの
DROP_KEYS = {"created_at", "updated_at", "exported_at", "reverted_at", "change_group"}
DATE_STAMP = re.compile(r"R\d{6}(?=_)")


class _Rollback(Exception):
    """atomic を巻き戻すためだけの例外"""


def plain(value):
    """JSON に書ける素の値へ（numpy / pandas / 日付 / NaN）"""
    if value is None:
        return None
    if hasattr(value, "item") and not isinstance(value, (str, bytes)):
        try:
            value = value.item()
        except (ValueError, AttributeError):
            pass
    if isinstance(value, float) and math.isnan(value):
        return None
    if type(value).__name__ in ("NAType", "NaTType"):
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if hasattr(value, "isoformat"):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, (set, frozenset)):
        return sorted(plain(v) for v in value)
    return value


def norm(obj, idmap):
    if isinstance(obj, dict):
        out = {}
        for k, v in obj.items():
            k = str(k)
            if k in DROP_KEYS:
                continue
            if k in TX_ID_KEYS:
                out[k] = idmap.get(plain(v), None if v is None else f"?{v}")
            elif k in TX_ID_LIST_KEYS:
                out[k] = [idmap.get(plain(x), f"?{x}") for x in v]
            elif k == "tx_ids_json":
                out["tx_ids"] = [idmap.get(x, f"?{x}") for x in json.loads(v)]
            elif k.endswith("_id"):
                continue
            else:
                out[k] = norm(v, idmap)
        return out
    if isinstance(obj, (list, tuple)):
        return [norm(v, idmap) for v in obj]
    return plain(obj)


def write_json(path: Path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(data, ensure_ascii=False, indent=1, sort_keys=False) + "\n",
        encoding="utf-8",
    )


def df_records(df):
    return [{k: plain(v) for k, v in row.items()} for row in df.to_dict(orient="records")]


def error_record(e: CsvImportError):
    return {"error": type(e).__name__, **{k: plain(v) for k, v in e.to_dict().items()}}


def tx_idmap(case):
    ids = case.transactions.order_by("id").values_list("id", flat=True)
    return {tx_id: n for n, tx_id in enumerate(ids, start=1)}


# ---------------------------------------------------------------------------
# 出力ファイル
# ---------------------------------------------------------------------------

def response_meta(resp):
    disp = resp.get("Content-Disposition", "")
    return {
        "status": resp.status_code,
        "content_type": resp.get("Content-Type"),
        "disposition": DATE_STAMP.sub("<DATE>", disp),
    }


def dump_csv_response(resp):
    raw = resp.content
    text = raw.decode("utf-8-sig")
    return {
        **response_meta(resp),
        "bom": raw.startswith(b"\xef\xbb\xbf"),
        "crlf": b"\r\n" in raw,
        "rows": list(csv.reader(io.StringIO(text))),
    }


def dump_xlsx_response(resp):
    meta = response_meta(resp)
    if resp.status_code != 200:
        return meta
    wb = load_workbook(io.BytesIO(resp.content))
    sheets = []
    for ws in wb.worksheets:
        formats, values = {}, []
        for row in ws.iter_rows():
            values.append([plain(c.value) for c in row])
            for c in row:
                if c.value is not None and c.number_format != "General":
                    formats[c.coordinate] = c.number_format
        sheets.append({
            "title": ws.title,
            "tab_color": ws.sheet_properties.tabColor.rgb if ws.sheet_properties.tabColor else None,
            "freeze_panes": ws.freeze_panes,
            "merged": sorted(str(r) for r in ws.merged_cells.ranges),
            "column_widths": {k: d.width for k, d in sorted(ws.column_dimensions.items()) if d.width},
            "number_formats": formats,
            "values": values,
        })
    return {**meta, "sheets": sheets}


# ---------------------------------------------------------------------------
# 案件1件ぶん
# ---------------------------------------------------------------------------

FILTER_VARIANTS = {
    "default": {},
    "keyword_halfwidth": {"keyword": "ﾌﾘｺﾐ"},
    "keyword_multi": {"keyword": "振替 本店"},
    "amount_out_range": {"amount_type": "out", "amount_min": "10,000", "amount_max": "300000"},
    "amount_both_min": {"amount_min": "500000"},
    "amount_both_max": {"amount_max": "1000"},
    "category_exclude": {"category": [UNCATEGORIZED], "category_mode": "exclude"},
    "date_range": {"date_from": "2024-04-15", "date_to": "2024-05-15"},
    "sort_amount_desc": {"sort": "-amount_out"},
}


def filter_state(**over):
    state = {
        "bank": [], "account": [], "category": [], "category_mode": "include",
        "keyword": "", "amount_min": "", "amount_max": "", "amount_type": "both",
        "date_from": "", "date_to": "", "sort": "", "per_page": "100",
        "transfer_category": [], "transfer_category_mode": "include",
    }
    state.update(over)
    return state


def tx_rows(qs, idmap):
    return norm(list(qs.values()), idmap)


def dump_analysis(case, idmap):
    out = {}
    for name, over in FILTER_VARIANTS.items():
        data = AnalysisService.get_analysis_data(case, filter_state(**over))
        if data.get("no_data"):
            out[name] = data
            continue
        data = dict(data)
        data["all_txs"] = [idmap[t] for t in data["all_txs"].values_list("id", flat=True)]
        data["flagged_txs"] = [idmap[t] for t in data["flagged_txs"].values_list("id", flat=True)]
        out[name] = norm(data, idmap)
    return out


def dump_unclassified(case, idmap):
    out = {}
    for keyword in ("", "振替", "ATM"):
        qs = case.transactions.with_account_info().filter(category=UNCATEGORIZED).order_by("date", "id")
        groups = AnalysisService.build_unclassified_groups(qs, keyword)
        out[keyword or "(none)"] = {
            **norm(groups, idmap),
            "group_suggestions": json.loads(
                AnalysisService.build_group_suggestions(groups["groups"], case)
            ),
        }
    return out


def dump_exports(client, case):
    pk = case.pk
    get = client.get
    q = {"keyword": "振", "amount_type": "out", "amount_min": "1000"}
    out = {
        "csv_all": dump_csv_response(get(reverse("export-csv", args=[pk, "all"]))),
        "csv_transfers": dump_csv_response(get(reverse("export-csv", args=[pk, "transfers"]))),
        "csv_flagged": dump_csv_response(get(reverse("export-csv", args=[pk, "flagged"]))),
        "csv_filtered": dump_csv_response(get(reverse("export-csv-filtered", args=[pk]), q)),
        "xlsx_by_category": dump_xlsx_response(get(reverse("export-xlsx-by-category", args=[pk]))),
        "xlsx_monthly_cashflow": dump_xlsx_response(get(reverse("export-monthly-cashflow-xlsx", args=[pk]))),
        "xlsx_passbook_inventory": dump_xlsx_response(get(reverse("export-passbook-inventory", args=[pk]))),
    }
    resp = get(reverse("export-json", args=[pk]))
    out["json"] = {**response_meta(resp), "body": norm(json.loads(resp.content), {})}
    return out


def dump_mutation(case, fn):
    """分類の書き込みを試して結果を記録し、巻き戻す"""
    result = {}
    try:
        with db_transaction.atomic():
            before = case.classification_changes.count()
            idmap = tx_idmap(case)
            result["return"] = plain(fn(case))
            result["transactions"] = [
                {"id": idmap[t["id"]], "category": t["category"], "classification_score": t["classification_score"]}
                for t in case.transactions.order_by("id").values("id", "category", "classification_score")
            ]
            changes = ClassificationChange.objects.filter(case=case).order_by("id")[before:]
            result["changes"] = norm(
                [{k: v for k, v in c.items() if k != "id"} for c in changes.values()], idmap
            )
            raise _Rollback
    except _Rollback:
        pass
    return result


def dump_case(case, out_dir: Path, client: Client):
    idmap = tx_idmap(case)
    write_json(out_dir / "case.json", norm({
        "name": case.name,
        "reference_date": case.reference_date,
        "custom_patterns": case.custom_patterns,
        "accounts": list(case.accounts.order_by("bank_name", "branch_name", "account_number").values()),
    }, idmap))
    write_json(out_dir / "transactions.json", tx_rows(case.transactions.order_by("id"), idmap))
    write_json(out_dir / "monthly_cashflow.json", norm(AnalysisService.get_monthly_cashflow(case), idmap))
    write_json(out_dir / "classification_preview.json",
               norm(TransactionService.get_classification_preview(case), idmap))
    write_json(out_dir / "analysis.json", dump_analysis(case, idmap))
    write_json(out_dir / "unclassified_groups.json", dump_unclassified(case, idmap))
    write_json(out_dir / "exports.json", dump_exports(client, case))
    write_json(out_dir / "apply_classification_rules.json",
               dump_mutation(case, TransactionService.apply_classification_rules))
    write_json(out_dir / "run_classifier.json", dump_mutation(case, TransactionService.run_classifier))


# ---------------------------------------------------------------------------
# 照合・分類の単体（DB を使わない）
# ---------------------------------------------------------------------------

def keyword_universe(extra_patterns=None):
    patterns = dict(config.get_classification_patterns())
    for cat, kws in (extra_patterns or {}).items():
        patterns[cat] = list(patterns.get(cat, [])) + list(kws)
    return sorted({kw for kws in patterns.values() for kw in kws})


def dump_fuzzy(texts, keywords):
    """rapidfuzz の点数そのもの。

    分類は摘要を**正規化せずに**そのまま rapidfuzz へ渡している（NFKC が掛かるのは検索用の
    description_search だけ）ので、点数も生の文字列どうしで取る。normalized は検索の正解として残す。
    """
    rows = []
    for text in texts:
        scores = {}
        for kw in keywords:
            scores[kw] = [round(fuzz.partial_ratio(text, kw), 4), round(fuzz.token_set_ratio(text, kw), 4)]
        rows.append({"text": text, "normalized": normalize_text(text), "scores": scores})
    return {"keywords": keywords, "score_order": ["partial_ratio", "token_set_ratio"], "rows": rows}


def dump_classify(texts, case_patterns=None):
    global_patterns = config.get_classification_patterns()
    fuzzy_config = config.get_fuzzy_config()
    rows = []
    for text in texts:
        row = {"text": text}
        for label, out_amt in (("out_0", 0), ("out_1000000", 1_000_000), ("out_999999", 999_999)):
            row[label] = list(llm_classifier.classify_by_rules(
                text, out_amt, 0,
                case_patterns=case_patterns, global_patterns=global_patterns, fuzzy_config=fuzzy_config,
            ))
        row["suggestions"] = [list(s) for s in llm_classifier.get_fuzzy_suggestions(
            text, case_patterns=case_patterns, global_patterns=global_patterns,
            fuzzy_config=fuzzy_config, top_n=3,
        )]
        rows.append(row)
    llm_classifier.clear_fuzzy_cache()
    return {"fuzzy_config": fuzzy_config, "case_patterns": case_patterns, "rows": rows}


def dump_importer(path: Path):
    content = path.read_bytes()
    out = {}
    for allow_multiple in (False, True):
        key = f"allow_multiple_{str(allow_multiple).lower()}"
        try:
            df = load_csv(io.BytesIO(content), allow_multiple=allow_multiple)
            attrs = {k: plain(v) if not isinstance(v, dict) else {kk: plain(vv) for kk, vv in v.items()}
                     for k, v in df.attrs.items()}
            checked = validate_balance(df)
            out[key] = {
                "columns": list(df.columns),
                "attrs": attrs,
                "rows": df_records(df),
                "validated": df_records(checked),
            }
        except CsvImportError as e:
            out[key] = error_record(e)
        except Exception as e:  # noqa: BLE001 ― 想定外の落ち方も「正解」として残す
            out[key] = {"unexpected": type(e).__name__, "message": str(e)}
    return out


# ---------------------------------------------------------------------------
# 取込ウィザード（画面と同じ POST を通す）
# ---------------------------------------------------------------------------

def wizard_parse(client, case, files: list[Path]):
    data = {"action": "parse_files"}
    for i, p in enumerate(files):
        data[f"file_{i}"] = SimpleUploadedFile(p.name, p.read_bytes())
    resp = client.post(reverse("import-wizard", args=[case.pk]), data,
                       HTTP_X_REQUESTED_WITH="XMLHttpRequest")
    return resp.status_code, json.loads(resp.content)


def _js_int(value):
    """`parseInt(String(v || 0)) || 0` と同じ"""
    m = re.match(r"\s*([+-]?\d+)", str(value or 0))
    return int(m.group(1)) if m else 0


def ui_commit_payload(parsed: dict, skip_duplicates: bool, typed_account: str = ""):
    """import_wizard.html の送信処理（submit ハンドラ）が組み立てるものと同じ形にする。

    行は画面の表から作り直される: 摘要の空欄は ''、残高の空欄は 0、1ファイルでも残高列が
    無ければ全ファイルの残高が null。送るのは `skipDuplicates` で、サーバが読む
    `duplicateAction` は送られない。口座番号が空のファイルがあると画面が先へ進ませないので、
    利用者が手で打った番号として typed_account を入れる（無ければ None を返す）。
    """
    files = parsed.get("files", [])
    no_balance = any(f.get("has_balance") is False for f in files)
    out = []
    for index, f in enumerate(files):
        det = f.get("detected_account") or {}
        account_number = str(det.get("account_number") or "").strip() or typed_account
        if not account_number:
            return None
        rows = []
        for r in f.get("rows", []):
            if skip_duplicates and r.get("is_duplicate"):
                continue
            rows.append({
                "date": r.get("date") or None,
                "description": r.get("description") or "",
                "amount_out": _js_int(r.get("amount_out")),
                "amount_in": _js_int(r.get("amount_in")),
                "balance": None if no_balance else _js_int(r.get("balance")),
                "is_new": False,
            })
        out.append({
            "filename": f.get("filename"),
            "rows": rows,
            "account": {
                "fileIndex": index,
                "bank_name": str(det.get("bank_name") or "").strip(),
                "branch_name": str(det.get("branch_name") or "").strip(),
                "account_type": str(det.get("account_type") or "").strip(),
                "account_number": account_number,
            },
        })
    return {"files": out, "skipDuplicates": skip_duplicates}


def wizard_commit(client, case, parsed: dict, skip_duplicates: bool, typed_account: str = ""):
    payload = ui_commit_payload(parsed, skip_duplicates, typed_account)
    if payload is None:
        return "blocked: 口座番号が空（画面が先へ進ませない）"
    resp = client.post(reverse("import-wizard", args=[case.pk]),
                       {"action": "commit_wizard", "wizard_data": json.dumps(payload, ensure_ascii=False)})
    return resp.status_code


class Command(BaseCommand):
    help = "React 移行の正解（いまの挙動）を JSON に書き出す"

    def add_arguments(self, parser):
        parser.add_argument("mode", choices=["synthetic", "real"])
        parser.add_argument("--inputs", help="synthetic: inputs ディレクトリ（files/ と scenarios.json）")
        parser.add_argument("--out", required=True)

    def handle(self, *args, mode, inputs, out, **options):
        out_dir = Path(out)
        client = Client()
        if mode == "synthetic":
            if not inputs:
                raise CommandError("--inputs が要ります")
            self.synthetic(Path(inputs), out_dir, client)
        else:
            self.real(out_dir, client)

    def synthetic(self, inputs: Path, out_dir: Path, client: Client):
        if Case.objects.exists():
            raise CommandError("空の DB に対して走らせてください（既存の案件があります）")
        files_dir = inputs / "files"
        files = sorted(p for p in files_dir.iterdir() if p.is_file())

        for p in files:
            write_json(out_dir / "importer" / f"{p.name}.json", dump_importer(p))
        self.stdout.write(f"importer: {len(files)} files")

        texts = [t for t in (inputs / "fuzzy_texts.txt").read_text(encoding="utf-8").splitlines()
                 if t and not t.startswith("#")]
        scenarios = json.loads((inputs / "scenarios.json").read_text(encoding="utf-8"))
        extra = {}
        for s in scenarios:
            for cat, kws in (s["case"].get("custom_patterns") or {}).items():
                extra.setdefault(cat, []).extend(kws)
        write_json(out_dir / "fuzzy.json", dump_fuzzy(texts, keyword_universe(extra)))
        write_json(out_dir / "classify.json", {
            "global_only": dump_classify(texts),
            "with_case_patterns": dump_classify(texts, scenarios[0]["case"].get("custom_patterns")),
        })

        for s in scenarios:
            self.run_scenario(s, files_dir, out_dir / "scenarios" / s["name"], client)

    def run_scenario(self, s, files_dir, sdir, client):
        c = s["case"]
        case = Case.objects.create(
            name=c["name"],
            reference_date=c.get("reference_date"),
            custom_patterns=c.get("custom_patterns") or {},
        )
        case.refresh_from_db()  # reference_date を文字列のまま持たせない
        steps = []
        for step_no, step in enumerate(s["steps"], start=1):
            if "parse" in step:
                status, parsed = wizard_parse(client, case, [files_dir / n for n in step["parse"]])
                rec = {"parse": step["parse"], "status": status, "response": norm(parsed, {})}
                if step.get("commit") and parsed.get("success"):
                    rec["commit"] = step["commit"]
                    # "skip" = 画面の「重複を除外」にチェック（既定）、"import" = 外した状態
                    # 口座番号を読み取れなかったファイルは、画面で手入力した想定の番号で取り込む
                    typed = f"9{step_no:06d}"
                    if any(not (f.get("detected_account") or {}).get("account_number")
                           for f in parsed.get("files", [])):
                        rec["typed_account"] = typed
                    rec["commit_status"] = wizard_commit(client, case, parsed,
                                                         skip_duplicates=step["commit"] == "skip",
                                                         typed_account=typed)
                    rec["transaction_count"] = case.transactions.count()
                steps.append(rec)
            if "flag" in step:
                idmap = {n: t for t, n in tx_idmap(case).items()}
                for n, memo in step["flag"].items():
                    TransactionService.toggle_flag(case, idmap[int(n)])
                    if memo:
                        TransactionService.update_memo(case, idmap[int(n)], memo)
                steps.append({"flag": step["flag"]})
        write_json(sdir / "steps.json", steps)
        dump_case(case, sdir, client)
        self.stdout.write(f"scenario {s['name']}: {case.transactions.count()} transactions")

    def real(self, out_dir: Path, client: Client):
        texts = set()
        for n, case in enumerate(Case.objects.order_by("id"), start=1):
            dump_case(case, out_dir / f"case_{n:02d}", client)
            texts.update(d for d in case.transactions.values_list("description", flat=True) if d)
            self.stdout.write(f"case_{n:02d}: {case.transactions.count()} transactions")
        texts = sorted(texts)
        write_json(out_dir / "fuzzy.json", dump_fuzzy(texts, keyword_universe()))
        write_json(out_dir / "classify.json", {"global_only": dump_classify(texts)})
